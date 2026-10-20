"""
tools/make_demo_data.py — 데모용 가상 데이터 생성 스크립트.

실제 데이터와 동일한 폴더 구조·파일 형식·컬럼 구조를 갖지만, 상품·채널·금액은
모두 임의로 만든 가상의 값이다. (실제 회사 데이터와 무관)

실행:
    python tools/make_demo_data.py            # ./demo_data 에 생성
    python tools/make_demo_data.py --out DIR  # 다른 위치에 생성

생성 구조 (utils/io/config.py의 build_paths()와 동일):
    demo_data/{YYYY}/{MM}/
        outputs/   vnb.parquet, vnb-walk.parquet
        dashboard/ trad.parquet, isp.parquet, var.parquet, commission.xlsx,
                   model-expense.csv, actual-expense.xlsx, lob-mapping.xlsx
        scenarios/ fa-spread-{YYYYMM}.xlsx, VNBQ{n}_INT_RETURN_MR.fac
"""
from __future__ import annotations

import argparse
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook

SEED = 20261026

# ── 분기 범위 ────────────────────────────────────────────────────
QUARTERS = [(2025, "Q4"), (2026, "Q1"), (2026, "Q2"), (2026, "Q3"), (2026, "Q4")]
Q_MONTH = {"Q1": "03", "Q2": "06", "Q3": "09", "Q4": "12"}
Q_END = {"Q1": "03-31", "Q2": "06-30", "Q3": "09-30", "Q4": "12-31"}
Q_MONTHS = {"Q1": [1, 2, 3], "Q2": [4, 5, 6], "Q3": [7, 8, 9], "Q4": [10, 11, 12]}

# 분기말 환율 (가상) — .env/config 기본값의 CURR/PREV FX와 맞춤
FX = {(2025, "Q4"): 1420.0, (2026, "Q1"): 1445.0, (2026, "Q2"): 1462.0,
      (2026, "Q3"): 1488.0, (2026, "Q4"): 1475.0}

# ── 가상 채널 ────────────────────────────────────────────────────
CHANNELS = ["AGENCY", "GA", "BANCA", "TM"]
CH_WEIGHT = {"AGENCY": 0.42, "GA": 0.30, "BANCA": 0.18, "TM": 0.10}
GA_TYPES = ["TYPE_A", "TYPE_B"]

STEP_ORDER = ["MC", "RP", "SL_VOL", "SL_MIX", "SL_DEMO", "AU_COMM",
              "EC_NMR", "EC_FX", "SL_NEW", "OTH_COMM", "OTH_ETU"]


@dataclass
class Product:
    code: str          # PRODUCT / PRODUCT_NAME_BASE
    item: str          # 표시용 이름(PRODUCT_ITEM)
    group: str         # PRODUCT_GROUP1
    mpf_file: str      # trad / isp / var
    plicd: str         # 5자리 가상 보종코드
    currency: str
    base_anp: float    # 분기 ANP (원) 기준값
    margin: float      # 기본 margin (VNB/ANP)
    tvog: float        # TVOG 비용 비율 (ANP 대비)
    terms: list
    channels: list
    defer: bool = False
    rate_sensitive: bool = False
    launch: tuple | None = None   # 출시 분기 (없으면 기존 상품)
    comm: float = 1.6             # 1회차+13회차 commission rate 기준


PRODUCTS = [
    Product("WL_CLASSIC", "Whole Life Classic", "Whole Life", "trad", "10101", "KRW", 9.0e9, 0.32, 0.000, [10, 15, 20], ["AGENCY", "GA", "TM"]),
    Product("WL_PLUS", "Whole Life Plus", "Whole Life", "trad", "10102", "KRW", 6.0e9, 0.28, 0.000, [10, 20], ["AGENCY", "GA", "BANCA"]),
    Product("WL_DOLLAR", "Dollar Whole Life", "Whole Life", "isp", "10201", "USD", 5.0e9, 0.22, 0.020, [5, 10], ["AGENCY", "GA", "BANCA"], rate_sensitive=True),
    Product("CI_CARE", "CI Care", "Health", "trad", "20101", "KRW", 7.0e9, 0.45, 0.000, [10, 20, 30], ["AGENCY", "GA", "TM"], comm=2.0),
    Product("CI_SIMPLE", "Simplified CI", "Health", "trad", "20102", "KRW", 4.0e9, 0.38, 0.000, [10, 20], ["AGENCY", "GA", "TM"], comm=1.9),
    Product("HEALTH_SENIOR", "Senior Health", "Health", "trad", "20103", "KRW", 3.0e9, 0.40, 0.000, [10, 20], ["AGENCY", "GA"], launch=(2026, "Q3"), comm=1.8),
    Product("TERM_BASIC", "Term Basic", "Term", "trad", "30101", "KRW", 2.0e9, 0.50, 0.000, [10, 20], ["AGENCY", "TM"], comm=1.5),
    Product("ANNUITY_SAVER", "Annuity Saver", "Annuity", "isp", "40101", "KRW", 6.0e9, 0.08, 0.030, [5, 7, 10], ["AGENCY", "BANCA"], defer=True, rate_sensitive=True, comm=0.9),
    Product("ANNUITY_DOLLAR", "Dollar Annuity", "Annuity", "isp", "40201", "USD", 3.0e9, 0.10, 0.040, [5, 10], ["AGENCY", "BANCA"], defer=True, rate_sensitive=True, comm=0.9),
    Product("VUL_GROWTH", "VUL Growth", "Variable", "var", "50101", "KRW", 4.0e9, 0.18, 0.060, [10, 20], ["AGENCY", "GA"], comm=1.4),
    Product("VA_INCOME", "VA Income", "Variable", "var", "50201", "KRW", 2.5e9, 0.12, 0.080, [5, 10], ["AGENCY", "BANCA"], defer=True, comm=1.0),
]


def _q_index(q):
    return QUARTERS.index(q)


def _launched(p: Product, q) -> bool:
    return p.launch is None or _q_index(q) >= _q_index(p.launch)


# ════════════════════════════════════════════════════════════════
# 1. 분기별 상태(state) — 채널×상품 셀 단위 ANP·VNB 구성요소
# ════════════════════════════════════════════════════════════════
def build_states(rng: np.random.Generator) -> tuple[dict, dict]:
    """
    state[q][(ch, code)] = {anp, core, etu, tvog, commtu}
      core  : EXPTU=0, TVOG=0 기준 VNB (commission true-up 포함)
      VNB(1,0) = core + etu,  VNB(1,1) = core + etu - tvog
    walk[q]  : 전분기 → 당분기 cumulative walk 행 목록
    """
    states: dict = {}
    walks: dict = {}

    # 분기별 공통 충격 (가상)
    growth = {q: g for q, g in zip(QUARTERS, [1.00, 0.97, 1.06, 1.04, 1.02])}
    nmr_shock = {q: s for q, s in zip(QUARTERS, [0.0, 0.004, -0.006, 0.008, -0.003])}
    au_comm = {q: s for q, s in zip(QUARTERS, [0.0, -0.002, 0.000, -0.004, 0.001])}

    for qi, q in enumerate(QUARTERS):
        st_q = {}
        for p in PRODUCTS:
            for ch in p.channels:
                key = (ch, p.code)
                if not _launched(p, q):
                    continue
                w = CH_WEIGHT[ch] / sum(CH_WEIGHT[c] for c in p.channels)
                if qi == 0 or key not in states.get(QUARTERS[qi - 1], {}):
                    anp = p.base_anp * w * rng.uniform(0.85, 1.15)
                    if p.launch == q:
                        anp *= 0.7
                else:
                    prev = states[QUARTERS[qi - 1]][key]
                    anp = prev["anp"] * growth[q] * rng.uniform(0.92, 1.10)
                etu = anp * rng.uniform(-0.012, 0.010)
                tvog = anp * p.tvog * rng.uniform(0.85, 1.20)
                commtu = anp * rng.uniform(-0.006, 0.006)
                st_q[key] = {"anp": anp, "etu": etu, "tvog": tvog, "commtu": commtu, "core": None}
        states[q] = st_q

        if qi == 0:
            for (ch, code), s in st_q.items():
                p = next(x for x in PRODUCTS if x.code == code)
                s["core"] = s["anp"] * p.margin * rng.uniform(0.9, 1.1) + s["commtu"]
            continue

        # ── walk: 전분기 상태 → 당분기 상태 ─────────────────────
        pq = QUARTERS[qi - 1]
        prev_st = states[pq]
        ratio = sum(s["anp"] for s in st_q.values()) / sum(s["anp"] for s in prev_st.values())
        fx_ratio = FX[q] / FX[pq]
        rows = []
        for (ch, code), s in st_q.items():
            p = next(x for x in PRODUCTS if x.code == code)
            ps = prev_st.get((ch, code))
            base = {"CHANNEL": ch, "PRODUCT_GROUP1": p.group, "PRODUCT": code}
            if ps is None:  # 신상품: SL_NEW에서 전액 반영
                core = s["anp"] * p.margin * rng.uniform(0.9, 1.1)
                cum = 0.0
                seq = [("PQ_COMM", 0.0, 0.0)] + [(st_, 0.0, 0.0) for st_ in STEP_ORDER[:8]]
                for st_, _, _ in seq:
                    rows.append({**base, "STEP": st_, "VNB": 0.0, "ANP": 0.0})
                cum = core - s["commtu"]
                rows.append({**base, "STEP": "SL_NEW", "VNB": cum, "ANP": s["anp"]})
                cum += s["commtu"]
                rows.append({**base, "STEP": "OTH_COMM", "VNB": cum, "ANP": s["anp"]})
                s["core"] = cum
                cum += s["etu"]
                rows.append({**base, "STEP": "OTH_ETU", "VNB": cum, "ANP": s["anp"]})
                continue

            pa, ca = ps["anp"], s["anp"]
            cum = ps["core"]
            cum -= ps["commtu"]
            rows.append({**base, "STEP": "PQ_COMM", "VNB": cum, "ANP": pa})
            cum += pa * rng.uniform(-0.004, 0.003)
            rows.append({**base, "STEP": "MC", "VNB": cum, "ANP": pa})
            cum += pa * rng.uniform(-0.002, 0.002)
            rows.append({**base, "STEP": "RP", "VNB": cum, "ANP": pa})
            m_base = cum / pa
            cum *= ratio
            rows.append({**base, "STEP": "SL_VOL", "VNB": cum, "ANP": pa * ratio})
            target = m_base * ca
            cum += 0.7 * (target - cum)
            rows.append({**base, "STEP": "SL_MIX", "VNB": cum, "ANP": ca})
            cum = target + ca * rng.uniform(-0.006, 0.006)
            rows.append({**base, "STEP": "SL_DEMO", "VNB": cum, "ANP": ca})
            cum += ca * au_comm[q] * p.comm
            rows.append({**base, "STEP": "AU_COMM", "VNB": cum, "ANP": ca})
            sens = 2.5 if p.rate_sensitive else 0.6
            cum += ca * nmr_shock[q] * sens
            rows.append({**base, "STEP": "EC_NMR", "VNB": cum, "ANP": ca})
            if p.currency == "USD":
                cum += cum * (fx_ratio - 1.0) * 0.5
            rows.append({**base, "STEP": "EC_FX", "VNB": cum, "ANP": ca})
            rows.append({**base, "STEP": "SL_NEW", "VNB": cum, "ANP": ca})
            cum += s["commtu"]
            rows.append({**base, "STEP": "OTH_COMM", "VNB": cum, "ANP": ca})
            s["core"] = cum
            cum += s["etu"]
            rows.append({**base, "STEP": "OTH_ETU", "VNB": cum, "ANP": ca})
        walks[q] = rows
    return states, walks


# ════════════════════════════════════════════════════════════════
# 2. 파일 작성
# ════════════════════════════════════════════════════════════════
def write_vnb(states, q, out: Path):
    date = f"{q[0]}-{Q_END[q[1]]}"
    prod = {p.code: p for p in PRODUCTS}
    rows = []
    for (ch, code), s in states[q].items():
        p = prod[code]
        common = {"DATE": date, "SCEN": "MR", "CHANNEL": ch, "PRODUCT_GROUP1": p.group,
                  "PRODUCT_GROUP2": p.mpf_file.upper(), "PRODUCT": code, "CURRENCY": p.currency,
                  "ANP": s["anp"]}
        rows.append({**common, "EXPTU_IDX": 0, "TVOG_IDX": 0, "VNB": s["core"]})
        rows.append({**common, "EXPTU_IDX": 1, "TVOG_IDX": 0, "VNB": s["core"] + s["etu"]})
        rows.append({**common, "EXPTU_IDX": 1, "TVOG_IDX": 1, "VNB": s["core"] + s["etu"] - s["tvog"]})
    pd.DataFrame(rows).to_parquet(out / "vnb.parquet", index=False)


def write_walk(walks, q, out: Path):
    pd.DataFrame(walks[q]).to_parquet(out / "vnb-walk.parquet", index=False)


def write_mpf(states, q, out: Path, rng: np.random.Generator):
    qi = _q_index(q)
    prod = {p.code: p for p in PRODUCTS}
    by_file: dict = {"trad": [], "isp": [], "var": []}
    age_shift = (qi - 2) * 0.6  # 분기가 지날수록 가입연령 소폭 상승 (가상)
    for (ch, code), s in states[q].items():
        p = prod[code]
        n = int(np.clip(s["anp"] / 2.5e7, 40, 220))
        weights = rng.gamma(2.0, 1.0, n)
        anp_krw = s["anp"] * weights / weights.sum()
        fx = FX[q] if p.currency == "USD" else 1.0
        anp = anp_krw / fx
        base_age = {"Annuity": 52, "Variable": 45, "Health": 41, "Term": 36, "Whole Life": 44}[p.group]
        ages = np.clip(rng.normal(base_age + age_shift, 9, n).round(), 20, 75).astype(int)
        terms = rng.choice(p.terms, n, p=None)
        defer = (rng.choice([0, 5, 10, 15, 20], n, p=[0.1, 0.3, 0.3, 0.2, 0.1]) if p.defer else np.zeros(n, int))
        ct = rng.choice(["Standard", "Simplified"], n, p=[0.7, 0.3] if p.group == "Health" else [0.9, 0.1])
        variant = rng.choice(["A", "B"], n, p=[0.6, 0.4])
        df = pd.DataFrame({
            "CHANNEL": ch,
            "PRODUCT": code,
            "PRODUCT_GROUP1": p.group,
            "PROD_NAME_MP": [f"{code}_{v}" for v in variant],
            "PROD_NAME_BASE": f"L{p.plicd}",
            "CONTRACT_TYPE": ct,
            "CURRENCY": p.currency,
            "AGE": ages,
            "PREM_PERIOD": terms,
            "PREM_TERM": terms,
            "DEFER_Y": defer,
            "ANP": anp,
            "ANNUAL_PREM_P2": anp,
            "SUM_ASSURED": anp * rng.uniform(8, 25, n),
            "CNT": 1,
        })
        by_file[p.mpf_file].append(df)
    for f, parts in by_file.items():
        if parts:
            pd.concat(parts, ignore_index=True).to_parquet(out / f"{f}.parquet", index=False)


def write_lob_mapping(out: Path):
    pd.DataFrame([{"PROD_NAME_BASE": f"L{p.plicd}", "PRODUCT": p.code, "PRODUCT_GROUP1": p.group}
                  for p in PRODUCTS]).to_excel(out / "lob-mapping.xlsx", index=False)


def write_commission(q, out: Path, rng: np.random.Generator):
    """{CHANNEL}_actual / {CHANNEL}_plan 시트 (row1/row2 이중 헤더, row3부터 데이터)."""
    wb = Workbook()
    wb.remove(wb.active)
    months = [f"{q[0]}-{m:02d}" for m in Q_MONTHS[q[1]]]
    qi = _q_index(q)
    for ch in CHANNELS:
        prods = [p for p in PRODUCTS if ch in p.channels and _launched(p, q)]
        if not prods:
            continue
        ga = ch == "GA"
        rows = []
        for p in prods:
            for term in p.terms:
                for gt in (GA_TYPES if ga else [None]):
                    seed_plan = p.comm * (1.08 if ga else 1.0) * (1 + 0.01 * term / 10)
                    plan1 = seed_plan * 0.78
                    plan13 = seed_plan * 0.22
                    drift = 1 + 0.015 * (qi - 2) + rng.uniform(-0.04, 0.04)
                    act = [(plan1 * drift * rng.uniform(0.97, 1.03), plan13 * drift * rng.uniform(0.97, 1.03))
                           for _ in months]
                    rows.append((p, term, gt, plan1, plan13, act))

        key_hdr = ["보종코드", "납기"] + (["GA_TYPE"] if ga else []) + ["보종명"]
        # actual
        ws = wb.create_sheet(f"{ch}_actual")
        top = ["PRODUCT_ITEM", "PRODUCT_NAME_BASE"] + [""] * len(key_hdr)
        sub = ["PRODUCT_ITEM", "PRODUCT_NAME_BASE"] + key_hdr
        for m in months:
            top += [m, m]
            sub += ["1회차", "13회차"]
        ws.append(top)
        ws.append(sub)
        for p, term, gt, _, _, act in rows:
            r = [p.item, p.code, p.plicd, str(term)] + ([gt] if ga else []) + [p.item]
            for a1, a13 in act:
                r += [round(a1, 4), round(a13, 4)]
            ws.append(r)
        # plan
        ws = wb.create_sheet(f"{ch}_plan")
        ws.append(["PRODUCT_ITEM", "PRODUCT_NAME_BASE"] + [""] * len(key_hdr) + ["PLAN", "PLAN"])
        ws.append(["PRODUCT_ITEM", "PRODUCT_NAME_BASE"] + key_hdr + ["1회차", "13회차"])
        for p, term, gt, plan1, plan13, _ in rows:
            ws.append([p.item, p.code, p.plicd, str(term)] + ([gt] if ga else []) + [p.item, round(plan1, 4), round(plan13, 4)])
    wb.save(out / "commission.xlsx")


_VARS = {"ANP": ("EEV_ACQ_EXP_ANP_IF", 0.09), "SC": ("EEV_ACQ_EXP_FYC_IF", 0.04),
         "PNI": ("EEV_ACQ_EXP_POL_IF", 0.03), "COMM": ("EEV_TOT_COMM", 0.55)}


def write_expense(states, q, out: Path, rng: np.random.Generator):
    nb_m = int(Q_MONTH[q[1]])
    model_rows, actual_rows = [], []
    totals = {d: 0.0 for d in _VARS}
    for ch in CHANNELS:
        ch_anp = sum(s["anp"] for (c, _), s in states[q].items() if c == ch)
        for drv, (var, ratio) in _VARS.items():
            model = ch_anp * ratio
            actual = model * rng.uniform(0.92, 1.12)
            model_rows.append({"LOB": "ALL", "NB_M": nb_m, "CHANNEL": ch, "VAR": var, "sum": model})
            actual_rows.append({"CHANNEL": ch, "DRIVER": drv, "EXPENSE": actual})
            totals[drv] += actual
    for drv, v in totals.items():
        actual_rows.append({"CHANNEL": "Company Total", "DRIVER": drv, "EXPENSE": v})
    pd.DataFrame(model_rows).to_csv(out / "model-expense.csv", index=False)
    pd.DataFrame(actual_rows).to_excel(out / "actual-expense.xlsx", index=False)


def write_scenarios(q, out: Path, rng: np.random.Generator):
    qi = _q_index(q)
    ym = f"{q[0]}{Q_MONTH[q[1]]}"
    spread = [{"PRODUCT": p.code, "SPREAD": round(0.55 + 0.25 * rng.random() + 0.03 * (qi - 2), 4)}
              for p in PRODUCTS if p.rate_sensitive]
    with pd.ExcelWriter(out / f"fa-spread-{ym}.xlsx", engine="openpyxl") as xw:
        pd.DataFrame(spread).to_excel(xw, sheet_name="Summary", index=False)

    shift = [0.0, 0.0008, -0.0006, 0.0012, -0.0004][qi]
    lines = ["VARIABLE_TYPES,T1,I,N,N", "!2,T,TL_GA_KRW,TL_GA_USD"]
    for t in range(1, 1201):  # 월 단위 1200개월 projection
        krw = 0.026 + 0.010 * (1 - np.exp(-t / 150)) + shift
        usd = 0.039 + 0.006 * (1 - np.exp(-t / 120)) + shift * 0.6
        lines.append(f"*,{t},{krw:.6f},{usd:.6f}")
    (out / f"VNBQ{q[1][1]}_INT_RETURN_MR.fac").write_text("\n".join(lines) + "\n", encoding="utf-8")


def main(out_root: Path) -> None:
    rng = np.random.default_rng(SEED)
    states, walks = build_states(rng)
    for q in QUARTERS:
        base = out_root / str(q[0]) / Q_MONTH[q[1]]
        outputs, dash, scen = base / "outputs", base / "dashboard", base / "scenarios"
        for d in (outputs, dash, scen):
            d.mkdir(parents=True, exist_ok=True)
        write_vnb(states, q, outputs)
        if q in walks:
            write_walk(walks, q, outputs)
        write_mpf(states, q, dash, rng)
        write_lob_mapping(dash)
        write_commission(q, dash, rng)
        write_expense(states, q, dash, rng)
        write_scenarios(q, scen, rng)
        print(f"generated {q[0]} {q[1]} → {base}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="navi 데모용 가상 데이터 생성")
    ap.add_argument("--out", default=str(Path(__file__).resolve().parent.parent / "demo_data"))
    main(Path(ap.parse_args().out))
